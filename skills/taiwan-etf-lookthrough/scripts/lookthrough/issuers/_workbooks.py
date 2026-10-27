"""Read bounded official workbooks without trusting formulas or dimensions."""

from io import BytesIO
from itertools import islice
from typing import Any
import warnings
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ._assets import MAX_RAW_BYTES


def read_sheets(raw: bytes) -> dict[str, list[list[Any]]]:
    """Read original XLSX cells with ZIP, row, column and formula limits.

    Some official XLSX exports declare A1 dimensions despite containing a
    complete table. Reset read-only dimensions, then apply our own limits.
    No macros, formulas or external links are evaluated. Return cell values
    keyed by worksheet name; raise ValueError for malformed/oversized input.
    """
    if not isinstance(raw, bytes) or not 0 < len(raw) <= MAX_RAW_BYTES:
        raise ValueError("Invalid workbook response size")
    try:
        with ZipFile(BytesIO(raw)) as archive:
            files = archive.infolist()
            if (len(files) > 100 or len({f.filename for f in files})
                    != len(files)
                    or sum(f.file_size for f in files) > 40 * 1024 * 1024
                    or any(f.flag_bits & 1 or f.file_size > MAX_RAW_BYTES
                           or "vbaproject" in f.filename.lower()
                           or "externallinks/" in f.filename.lower()
                           for f in files)):
                raise ValueError("Unsafe or oversized source workbook")
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="Workbook contains no default style.*",
                category=UserWarning)
            book = load_workbook(BytesIO(raw), read_only=True,
                                 data_only=False, keep_links=False)
        try:
            if not 1 <= len(book.worksheets) <= 12:
                raise ValueError("Unknown source workbook sheet count")
            result = {}
            for sheet in book:
                sheet.reset_dimensions()
                rows = []
                for row in islice(sheet.iter_rows(), 6001):
                    if (len(row) > 12 or any(
                            cell.data_type == "f"
                            or len(str(cell.value)) > 10000 for cell in row)):
                        raise ValueError("Unqualified workbook cell/layout")
                    values = [cell.value for cell in row]
                    while values and values[-1] is None:
                        values.pop()
                    rows.append(values)
                if len(rows) > 6000:
                    raise ValueError("Source workbook exceeds row limit")
                result[sheet.title] = rows
            return result
        finally:
            book.close()
    except (BadZipFile, InvalidFileException, ParseError, KeyError,
            EOFError, OSError) as error:
        raise ValueError("Malformed source workbook") from error
