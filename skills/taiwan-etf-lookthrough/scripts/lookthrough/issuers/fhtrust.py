"""Parse paired Fuh Hwa full XLSX and same-date official asset-page bytes."""

import base64
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from html.parser import HTMLParser
from io import BytesIO
import json
import re
from typing import Any, TYPE_CHECKING
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ._assets import (
    MAX_RAW_BYTES, check_percentage, check_stock, nav_number, number,
    read_json,
)
from ..source_quality import discrepancy
from ..models import Holding, Snapshot, snapshot_from_dict, snapshot_to_dict

if TYPE_CHECKING:
    from .registry import SupportedEtf


SOURCE_URL = "https://www.fhtrust.com.tw/ETF/etf_detail/ETF21"
_HOST = "https://www.fhtrust.com.tw"
_HEADER = ["證券代號", "證券名稱", "股數", "金額", "權重(%)"]
_FUTURES = ["期貨代號", "期貨名稱", "口數", "金額", "權重(%)"]
_ASSETS = {"扣除應付買入證券款後現金餘額(NTD)": "cash",
           "期貨保證金(NTD)": "other", "應付申購預收款(NTD)": "other"}


class _Page(HTMLParser):
    """Read only the official assets panel without evaluating HTML/JS."""

    def __init__(self, raw: bytes) -> None:
        """Collect the single panel, exact canonical URL, tables and link."""
        super().__init__(convert_charrefs=True)
        if not isinstance(raw, bytes) or not 0 < len(raw) <= MAX_RAW_BYTES:
            raise ValueError("Invalid FH page size")
        self.depth = 0
        self.panel_depth = -1
        self.panels = 0
        self.canonical: list[str] = []
        self.links: list[str] = []
        self.text: list[str] = []
        self.tables: list[list[list[str]]] = []
        self.table: list[list[str]] | None = None
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.feed(raw.decode("utf-8-sig"))
        self.close()
        if (self.panels != 1 or self.panel_depth != -1
                or self.canonical != [SOURCE_URL] or self.table is not None):
            raise ValueError("Malformed FH assets panel or page identity")

    def handle_starttag(self, tag: str,
                        attrs: list[tuple[str, str | None]]) -> None:
        """Retain source table boundaries and the asset download link."""
        if len(dict(attrs)) != len(attrs):
            raise ValueError("Duplicate FH HTML attribute")
        values = dict(attrs)
        if tag == "meta" and values.get("property") == "og:url":
            self.canonical.append(values.get("content", ""))
        if tag == "div":
            self.depth += 1
            if values.get("id") == "etfPanel3":
                self.panel_depth = self.depth
                self.panels += 1
        if self.panel_depth == -1:
            return
        if tag == "a" and "assetsExcel" in str(values.get("href")):
            self.links.append(values["href"])
        if tag == "table":
            if self.table is not None:
                raise ValueError("Nested FH assets table")
            self.table = []
        elif tag == "tr" and self.table is not None:
            if self.row is not None:
                raise ValueError("Unclosed FH table row")
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            if self.cell is not None:
                raise ValueError("Unclosed FH table cell")
            self.cell = []

    def handle_endtag(self, tag: str) -> None:
        """Finish exact rows; malformed closures are never auto-repaired."""
        if tag in ("td", "th") and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.cell is not None:
                raise ValueError("Malformed FH table cell")
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            if self.row is not None:
                raise ValueError("Malformed FH table row")
            self.tables.append(self.table)
            self.table = None
        if tag == "div":
            if self.depth == self.panel_depth:
                self.panel_depth = -1
            self.depth -= 1

    def handle_data(self, data: str) -> None:
        """Keep plain text only; never execute source scripts."""
        if self.panel_depth != -1:
            self.text.append(data)
            if self.cell is not None:
                self.cell.append(data)


@dataclass(frozen=True)
class AssetsPage:
    """Keep independently disclosed date, NAV, aggregate and assets."""

    day: date
    download_url: str
    nav: Decimal
    stock_total: Decimal
    futures_total: Decimal
    assets: tuple[Holding, ...]


def inspect_page(raw: bytes) -> AssetsPage:
    """Validate a current official panel and derive its dated XLSX URL."""
    page = _Page(raw)
    if len(page.links) != 1 or not re.fullmatch(
            r"/api/assetsExcel/ETF21/\d{8}", page.links[0]):
        raise ValueError("Unqualified FH asset download URL")
    stamp = page.links[0].rsplit("/", 1)[1]
    day = date.fromisoformat(f"{stamp[:4]}-{stamp[4:6]}-{stamp[6:]}")
    text = " ".join(" ".join(page.text).split())
    dates = re.findall(r"\(日期：(\d{4}/\d{2}/\d{2})\)", text)
    navs = re.findall(r"基金資產淨值 NTD ([\d,]+)", text)
    if dates != [day.strftime("%Y/%m/%d")] or len(navs) != 1:
        raise ValueError("Missing/ambiguous FH date or NAV")
    nav = nav_number(navs[0])
    summary = [t for t in page.tables if t and t[0]
               == ["項目", "金額", "權重 (%)"]]
    if len(summary) != 2:
        raise ValueError("Missing FH complete asset summaries")
    stock_total = None
    futures_total = None
    assets: list[Holding] = []
    seen: set[str] = set()
    for table in summary:
        for row in table[1:]:
            if len(row) != 3 or row[0] in seen:
                raise ValueError("Malformed/duplicate FH asset summary")
            label = row[0]
            seen.add(label)
            amount = number(row[1])
            if label == "股票":
                stock_total = amount
            elif label == "期貨":
                futures_total = amount
                continue
            elif label in _ASSETS and -nav <= amount <= nav:
                assets.append(Holding(label, label, _ASSETS[label],
                                      float(amount / nav)))
            else:
                raise ValueError("Unqualified FH asset label or amount")
    if (seen != {"股票", "期貨", *_ASSETS} or stock_total is None
            or futures_total is None):
        raise ValueError("Incomplete FH asset summaries")
    return AssetsPage(
        day, _HOST + page.links[0], nav, stock_total, futures_total,
        tuple(assets))


def capture(page: bytes, workbook: bytes) -> bytes:
    """Pair immutable original bytes for the standard three-argument parser.

    The returned envelope is a local capture, not an official JSON file.
    Original page/XLSX bytes and their hashes belong in source provenance.
    """
    if (not isinstance(workbook, bytes)
            or not 0 < len(workbook) <= MAX_RAW_BYTES):
        raise ValueError("Invalid FH XLSX size")
    inspect_page(page)
    return json.dumps({
        "format": "fhtrust-assets-v1", "page_html": page.decode("utf-8-sig"),
        "workbook_base64": base64.b64encode(workbook).decode("ascii"),
    }, ensure_ascii=False).encode("utf-8")


def _workbook_rows(raw: bytes) -> list[list[Any]]:
    """Read a bounded formula-free worksheet, rejecting unsafe ZIP growth."""
    try:
        with ZipFile(BytesIO(raw)) as archive:
            files = archive.infolist()
            if (len(files) > 100 or len({f.filename for f in files})
                    != len(files)
                    or sum(f.file_size for f in files) > 40 * 1024 * 1024
                    or any(f.flag_bits & 1 or f.file_size > MAX_RAW_BYTES
                           or "vbaProject" in f.filename
                           or "externalLinks/" in f.filename
                           for f in files)):
                raise ValueError("Unsafe or oversized FH XLSX archive")
        workbook = load_workbook(BytesIO(raw), read_only=True,
                                 data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("Unknown FH worksheet layout")
            sheet = workbook.active
            if sheet.max_row > 5200 or sheet.max_column != 5:
                raise ValueError("FH worksheet exceeds layout limits")
            rows = []
            for row in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Formula cells are not source amounts")
                rows.append([cell.value for cell in row])
            return rows
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, EOFError,
            ParseError) as error:
        raise ValueError("Malformed FH XLSX") from error


def parse(raw: bytes, etf: "SupportedEtf", fetched_at: str) -> Snapshot:
    """Parse complete paired sources, retaining signed asset contributions.

    All stock rows must reconcile to the independent page stock total and
    match amount/NAV at the XLSX's three-decimal percentage precision.
    Futures displayed weights are not interpreted as NAV or notional.
    No network access, file execution or source mutation occurs.
    """
    if (etf.symbol != "00929" or etf.issuer != "fhtrust" or etf.is_active
            or etf.source_url != SOURCE_URL):
        raise ValueError("Unqualified FH request identity")
    payload = read_json(raw, 24 * 1024 * 1024)
    if (not isinstance(payload, dict) or set(payload)
            != {"format", "page_html", "workbook_base64"}
            or payload["format"] != "fhtrust-assets-v1"
            or not isinstance(payload["page_html"], str)
            or not isinstance(payload["workbook_base64"], str)):
        raise ValueError("Invalid FH paired capture")
    page = inspect_page(payload["page_html"].encode("utf-8"))
    try:
        body = base64.b64decode(payload["workbook_base64"], validate=True)
    except ValueError as error:
        raise ValueError("Invalid FH workbook encoding") from error
    if not 0 < len(body) <= MAX_RAW_BYTES:
        raise ValueError("Invalid FH workbook size")
    rows = _workbook_rows(body)
    if (len(rows) < 15 or not isinstance(rows[0][0], str)
            or not rows[0][0].startswith("復華台灣科技優息ETF基金")
            or "（證劵代碼：00929）" not in rows[0][0]
            or rows[2][0] != f"日期: {page.day:%Y/%m/%d}"
            or rows[3][0] != "基金資產淨值"
            or nav_number(rows[4][0]) != page.nav or rows[10] != _HEADER):
        raise ValueError("FH workbook identity/date/NAV/header mismatch")
    holdings = list(page.assets)
    total = Decimal(0)
    futures_total = Decimal(0)
    seen: set[str] = set()
    futures = False
    for row in rows[11:]:
        if not any(value is not None for value in row):
            continue
        if row == _FUTURES and not futures:
            futures = True
            continue
        code, name = row[:2]
        if futures:
            if (not isinstance(code, str) or not code or code in seen
                    or not isinstance(name, str) or not name):
                raise ValueError("Malformed FH futures identity")
            quantity = number(row[2])
            if quantity < 0 or quantity != quantity.to_integral():
                raise ValueError("Invalid FH futures quantity")
            futures_total += number(row[3])
            if not isinstance(row[4], str) or not re.fullmatch(
                    r"-?\d+\.\d{3}%", row[4]):
                raise ValueError("Invalid FH futures percentage")
            holdings.append(Holding(code, name, "futures", None))
        else:
            amount = number(row[3])
            check_stock(code, name, number(row[2]), amount, page.nav)
            if code in seen:
                raise ValueError("Duplicate FH stock identity")
            check_percentage(row[4], amount, page.nav, 3)
            total += amount
            holdings.append(Holding(code, name, "stock",
                                    float(amount / page.nav)))
        seen.add(code)
    if total != page.stock_total or futures_total != page.futures_total:
        discrepancy("Complete FH stock/futures totals do not reconcile")
    snapshot = Snapshot(
        "1", etf.symbol, etf.issuer, False, page.day, fetched_at,
        page.download_url, "agent_official", tuple(holdings))
    return snapshot_from_dict(snapshot_to_dict(snapshot))
