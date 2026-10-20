"""Test source normalization and workbook boundaries using synthetic data."""

from dataclasses import asdict, replace
from datetime import date
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from taiwan_dcf import (  # noqa: E402
    Dataset, Fact, annual_value, collect_current_data,
    create_researched_workbook, inspect_ixbrl, load_supplement,
    parse_ixbrl, ttm_value, write_workbook,
)


FIXTURE = Path(__file__).parent / "fixtures" / "synthetic.json"
STAMP = "2026-10-20T00:00:00+00:00"


def fact(value: float, period: str, start: str, end: str) -> Fact:
    """Build a synthetic fact with explicit period and normalized units."""
    return Fact("revenue", value, start, end, period,
                "synthetic://unit-test", STAMP)


def document(
    value: str = "1,200", scale: str = "3", sign: str = "",
    company: str = "2330", category: str = "Consolidated report",
    currency: str = "iso4217:TWD", dimensional: bool = False,
    concept: str = "ifrs-full:Revenue",
) -> bytes:
    """Return minimal synthetic inline XBRL with real unit/context syntax."""
    dimension = ("<xbrldi:explicitMember>segment</xbrldi:explicitMember>"
                 if dimensional else "")
    return f"""<html>
    <ix:nonNumeric name="tifrs-notes:CompanyID">{company}</ix:nonNumeric>
    <ix:nonNumeric name="tifrs-notes:ReportCategory">{category}</ix:nonNumeric>
    <xbrli:context id="annual"><xbrli:entity>
      <xbrli:identifier>{company}</xbrli:identifier>{dimension}
    </xbrli:entity><xbrli:period><xbrli:startDate>2025-01-01</xbrli:startDate>
    <xbrli:endDate>2025-12-31</xbrli:endDate></xbrli:period></xbrli:context>
    <xbrli:unit id="money"><xbrli:measure>{currency}</xbrli:measure>
    </xbrli:unit><ix:nonFraction name="{concept}" contextRef="annual"
    unitRef="money" format="ixt:numdotdecimal" scale="{scale}"
    sign="{sign}">{value}</ix:nonFraction></html>""".encode()


class PeriodTests(unittest.TestCase):
    """Protect annual/YTD/quarter semantics and conflict handling."""

    def test_annual_wins_without_double_counting(self) -> None:
        """An annual report is not summed with its overlapping YTD rows."""
        rows = [fact(100, "annual", "2025-01-01", "2025-12-31"),
                fact(20, "ytd", "2025-01-01", "2025-03-31"),
                fact(45, "ytd", "2025-01-01", "2025-06-30")]
        self.assertEqual(annual_value(rows, "revenue", 2025), 100)

    def test_ytd_is_differenced_before_annual_aggregation(self) -> None:
        """Four cumulative values reconstruct one annual flow, not 150."""
        rows = [fact(value, "ytd", "2025-01-01", f"2025-{end}")
                for value, end in zip((10, 25, 45, 70),
                                      ("03-31", "06-30", "09-30", "12-31"))]
        self.assertEqual(annual_value(rows, "revenue", 2025), 70)
        self.assertIsNone(annual_value(rows[:-1], "revenue", 2025))

    def test_four_single_quarters_and_real_zero(self) -> None:
        """Four explicit quarters preserve a true numeric zero."""
        rows = [fact(value, "quarter", f"2025-{start}", f"2025-{end}")
                for value, start, end in (
                    (0, "01-01", "03-31"), (25, "04-01", "06-30"),
                    (30, "07-01", "09-30"), (35, "10-01", "12-31"))]
        self.assertEqual(annual_value(rows, "revenue", 2025), 90)
        self.assertIsNone(annual_value([], "revenue", 2025))

    def test_conflicting_duplicate_is_rejected(self) -> None:
        """Conflicting same-period facts must not become last-write-wins."""
        original = fact(100, "annual", "2025-01-01", "2025-12-31")
        with self.assertRaises(ValueError):
            annual_value([original, replace(original, value=101)],
                         "revenue", 2025)

    def test_mixed_quarter_and_ytd(self) -> None:
        """Known quarter totals can support YTD, but not inconsistent YTD."""
        rows = [fact(10, "quarter", "2025-01-01", "2025-03-31"),
                fact(25, "ytd", "2025-01-01", "2025-06-30"),
                fact(20, "quarter", "2025-07-01", "2025-09-30"),
                fact(70, "ytd", "2025-01-01", "2025-12-31")]
        self.assertEqual(annual_value(rows, "revenue", 2025), 70)

    def test_ttm_requires_matching_comparatives(self) -> None:
        """Subtract prior YTD once and never annualize missing data."""
        rows = [
            fact(100, "annual", "2025-01-01", "2025-12-31"),
            fact(40, "ytd", "2025-01-01", "2025-06-30"),
            fact(60, "ytd", "2026-01-01", "2026-06-30"),
        ]
        self.assertEqual(ttm_value(rows, "revenue", "2026-06-30"), 120)
        self.assertIsNone(ttm_value(rows[::2], "revenue", "2026-06-30"))
        with self.assertRaises(ValueError):
            ttm_value(rows, "cash", "2026-06-30")


class InlineXbrlTests(unittest.TestCase):
    """Check exact tags, scale, sign, scope, context and currency."""

    def test_scale_and_negative_cash_capex(self) -> None:
        """Displayed thousands become million once, and outflows positive."""
        rows = parse_ixbrl(document(), "2330", 2025, "synthetic://x", STAMP)
        self.assertEqual(rows[0].value, 1.2)
        self.assertEqual(rows[0].scale, 3)
        self.assertEqual(rows[0].published_at, "")
        rows = parse_ixbrl(document(sign="-", concept=(
            "ifrs-full:PurchaseOfPropertyPlantAndEquipmentClassifiedAs"
            "InvestingActivities")), "2330", 2025, "synthetic://x", STAMP)
        self.assertEqual(rows[0].value, 1.2)
        self.assertEqual(rows[0].metric, "capex_ppe")
        self.assertEqual(rows[0].sign, "-")

    def test_wrong_company_or_parent_scope_is_rejected(self) -> None:
        """A superficially valid report cannot cross entity or scope."""
        for body in (document(company="9999"),
                     document(category="Separate report")):
            with self.assertRaises(ValueError):
                parse_ixbrl(body, "2330", 2025, "synthetic://x", STAMP)

    def test_foreign_dimension_unknown_and_nonfinite_are_not_guessed(
        self,
    ) -> None:
        """Unsupported contexts and values do not silently enter a model."""
        for body in (document(currency="iso4217:USD"),
                     document(dimensional=True), document(value="NaN"),
                     document().replace(
                         b"<xbrli:identifier>2330",
                         b"<xbrli:identifier>9999"),
                     document(concept="custom:OperatingProfit")):
            self.assertEqual(parse_ixbrl(
                body, "2330", 2025, "synthetic://x", STAMP), [])

    def test_unmapped_concepts_remain_available_for_research(self) -> None:
        """An unknown tag can be inspected without silently mapping it."""
        body = document(concept="custom:UnmappedCashUse")
        rows = inspect_ixbrl(body, "2330", 2025)
        self.assertEqual(rows[0]["tag"], "custom:UnmappedCashUse")
        self.assertEqual(rows[0]["value"], 1.2)
        self.assertEqual(
            parse_ixbrl(body, "2330", 2025, "synthetic://x", STAMP), [])

    def test_roster_issued_shares_are_separate_from_outstanding(self) -> None:
        """Both roster schemas retain shares despite missing filings."""
        for key in ("IssueShares", "已發行普通股數或TDR原股發行股數"):
            identity = {"公司代號": "2330", "公司名稱": "SYNTHETIC",
                        "產業別": "24", key: "1,234,000"}
            with patch("taiwan_dcf._roster", return_value=(
                    [identity], "synthetic://roster")), patch(
                        "taiwan_dcf._fetch", return_value=b"[]"):
                data = collect_current_data("2330")
            issued = [row for row in data.facts
                      if row.metric == "issued_shares"]
            self.assertEqual(len(issued), 1)
            self.assertEqual(issued[0].value, 1.234)
            self.assertEqual(issued[0].period, "snapshot")
            self.assertFalse(any(row.metric == "shares"
                                 for row in data.facts))


class WorkbookTests(unittest.TestCase):
    """Verify formula artifacts; native-engine recalcs are separate."""

    def setUp(self) -> None:
        """Load the explicitly synthetic full five-year input fixture."""
        self.supplement = load_supplement(FIXTURE)
        self.data = Dataset("9999", "SYNTHETIC TEST COMPANY", "TWSE", "24",
                            date(2026, 10, 20))

    def test_complete_workbook_retains_formulas_and_sources(self) -> None:
        """One build references editable active case rows and source facts."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.xlsx"
            result = write_workbook(self.data, path, self.supplement)
            book = load_workbook(path)
            self.assertEqual(book.sheetnames,
                             ["Model", "Assumptions", "History", "Facts",
                              "Evidence"])
            self.assertEqual(book["History"]["G4"].value, 1000)
            self.assertIsNone(book["History"]["C23"].value)
            self.assertIn("D23", book["History"]["D22"].value)
            self.assertEqual(book["Assumptions"]["C4"].value, 2)
            self.assertIn("MAX(C22,0)", book["Model"]["C23"].value)
            self.assertIn("C53", book["Assumptions"]["C56"].value)
            self.assertIn("INDEX", book["Assumptions"]["C16"].value)
            self.assertIn("Assumptions!$C$11", book["Model"]["E50"].value)
            self.assertEqual(result["annual_metric_counts"]["2025"], 13)
            self.assertFalse(result["recalculated"])
            self.assertEqual(book.calculation.calcMode, "auto")
            self.assertEqual(book["Facts"]["F2"].value,
                             "synthetic://fixture")

    def test_missing_does_not_become_zero_or_default_wacc(self) -> None:
        """Incomplete data produces a workbook with blank required inputs."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "incomplete.xlsx"
            result = write_workbook(self.data, path)
            book = load_workbook(path)
            self.assertIsNone(book["Assumptions"]["C10"].value)
            self.assertIsNone(book["Assumptions"]["C12"].value)
            self.assertIsNone(book["Assumptions"]["E49"].value)
            self.assertIsNone(book["Assumptions"]["E57"].value)
            self.assertIsNone(book["History"]["C4"].value)
            self.assertIn('COUNT', book["Model"]["C42"].value)
            self.assertEqual(result["fact_count"], 0)
            with self.assertRaises(FileExistsError):
                write_workbook(self.data, path)

    def test_browser_research_builds_without_network_or_office(self) -> None:
        """Browser facts and assumptions reach Excel with no network calls."""
        metadata = asdict(self.data)
        metadata["valuation_date"] = self.data.valuation_date.isoformat()
        metadata["gaps"] = ["SYNTHETIC: Python network unavailable"]
        metadata["facts"] = self.supplement.pop("facts")
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            normalized = folder / "normalized.json"
            supplement = folder / "research.json"
            output = folder / "offline.xlsx"
            normalized.write_text(json.dumps(metadata), encoding="utf-8")
            supplement.write_text(json.dumps(self.supplement),
                                  encoding="utf-8")
            with (
                patch("taiwan_dcf.urlopen",
                      side_effect=AssertionError("Network disabled")) as net,
                patch("taiwan_dcf.recalculate_workbook",
                      side_effect=AssertionError("Office unavailable")) as rec,
            ):
                result = create_researched_workbook(
                    normalized, supplement, output)
            net.assert_not_called()
            rec.assert_not_called()
            book = load_workbook(output)
            self.assertEqual(book["History"]["G4"].value, 1000)
            self.assertEqual(book["Assumptions"]["C10"].value,
                             self.supplement["inputs"]["debt"]["value"])
            self.assertEqual(book["Model"]["C42"].data_type, "f")
            self.assertEqual(book["Facts"]["F2"].value,
                             "synthetic://fixture")
            self.assertEqual(result["annual_metric_counts"]["2025"], 13)
            self.assertFalse(result["recalculated"])

    def test_financial_company_is_not_applicable(self) -> None:
        """Even full inputs do not enable a financial-company FCFF result."""
        with tempfile.TemporaryDirectory() as directory:
            self.data.financial = True
            path = Path(directory) / "financial.xlsx"
            write_workbook(self.data, path, self.supplement)
            self.assertEqual(load_workbook(path)["Model"]["C16"].value, 0)

    def test_source_text_cannot_inject_formulas(self) -> None:
        """Untrusted source descriptions are stored as text cells."""
        self.supplement["inputs"]["debt"]["source"] = "=HYPERLINK(1)"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "safe.xlsx"
            write_workbook(self.data, path, self.supplement)
            cell = load_workbook(path)["Assumptions"]["I10"]
            self.assertEqual(cell.data_type, "s")

    def test_supplement_requires_provenance_and_finite_numbers(self) -> None:
        """Nonfinite numbers and text cannot become numeric inputs."""
        for value in (float("nan"), float("inf"), True, "0"):
            data = {"inputs": {"debt": {
                "value": value, "source": "synthetic://test", "method": "x"}}}
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "invalid.json"
                path.write_text(json.dumps(data), encoding="utf-8")
                with self.assertRaises(ValueError):
                    load_supplement(path)

    def test_current_period_conflict_and_invalid_period_are_rejected(
        self,
    ) -> None:
        """Check every period, including the latest balance sheet."""
        current = Fact("cash", 50, "", "2026-06-30", "instant",
                       "synthetic://test", STAMP)
        for rows in ([current, replace(current, value=51)],
                     [replace(current, period="annual")]):
            self.data.facts = rows
            with tempfile.TemporaryDirectory() as directory:
                with self.assertRaises(ValueError):
                    write_workbook(self.data, Path(directory) / "invalid.xlsx")

    def test_supplied_sources_survive_explanatory_notes(self) -> None:
        """Preserve terminal provenance in all three selectable cases."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sources.xlsx"
            write_workbook(self.data, path, self.supplement)
            sheet = load_workbook(path)["Assumptions"]
            for address in ("I11", "I57", "J57", "K57",
                            "I58", "J58", "K58"):
                self.assertIn("synthetic://fixture", sheet[address].value)

    def test_component_reconciliation_rejects_mismatched_sum(self) -> None:
        """An input total cannot silently disagree with its components."""
        debt = self.supplement["inputs"]["debt"]
        debt["components"] = [
            {"label": "Loan", "value": debt["value"] + 1, "sign": 1,
             "source": "synthetic://fixture", "method": "Test mismatch",
             "classification": "analyst_estimate"},
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mismatch.json"
            path.write_text(json.dumps(self.supplement))
            with self.assertRaisesRegex(ValueError, "Component sum"):
                load_supplement(path)

    def test_date_difference_is_disclosed_not_rejected(self) -> None:
        """An explicitly dated carry-forward proxy may still be used."""
        self.supplement["inputs"]["base_revenue"]["period_end"] = "2025-12-31"
        self.supplement["inputs"]["base_nwc"]["as_of"] = "2026-06-30"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "dates.json"
            path.write_text(json.dumps(self.supplement))
            accepted = load_supplement(path)
            self.assertEqual(
                accepted["inputs"]["base_nwc"]["as_of"], "2026-06-30")


if __name__ == "__main__":
    unittest.main()
