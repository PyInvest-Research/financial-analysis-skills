"""Collect official Taiwan filings and write an editable, gated FCFF model.

Amounts are TWD million, share counts are million shares, and prices are
TWD per share. This module writes Excel formulas; it does not evaluate DCF.
"""

from __future__ import annotations

import csv
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
import hashlib
from html.parser import HTMLParser
import io
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.properties import CalcProperties


TWSE = "https://openapi.twse.com.tw/v1/opendata/"
TPEX = "https://www.tpex.org.tw/openapi/v1/"
MOPS = "https://mopsov.twse.com.tw/server-java/t164sb01"
CASES = ("bear", "base", "bull")
CLASSIFICATIONS = ("reported", "derived", "analyst_estimate")
HISTORY_METRICS = (
    "revenue", "ebit", "pretax_income", "tax_expense", "depreciation",
    "amortisation", "capex_ppe", "capex_intangible", "nwc", "cash",
    "debt", "shares", "equity_adjustment",
)
TAGS = {
    "ifrs-full:Revenue": "revenue",
    "ifrs-full:ProfitLossFromOperatingActivities": "ebit",
    "ifrs-full:ProfitLossBeforeTax": "pretax_income",
    "ifrs-full:IncomeTaxExpenseContinuingOperations": "tax_expense",
    "ifrs-full:AdjustmentsForDepreciationExpense": "depreciation",
    "ifrs-full:AdjustmentsForAmortisationExpense": "amortisation",
    "ifrs-full:CashAndCashEquivalents": "cash",
    "ifrs-full:PurchaseOfPropertyPlantAndEquipmentClassifiedAsInvesting"
    "Activities": "capex_ppe",
    "ifrs-full:PurchaseOfIntangibleAssetsClassifiedAsInvestingActivities":
        "capex_intangible",
    "tifrs-bsci-ci:AccountsReceivableNet": "trade_receivables",
    "tifrs-bsci-ci:AccountsReceivableDuefromRelatedPartiesNet":
        "related_trade_receivables",
    "ifrs-full:OtherCurrentReceivables": "other_current_receivables",
    "tifrs-bsci-ci:OtherReceivablesDueFromRelatedParties":
        "related_other_receivables",
    "ifrs-full:Inventories": "inventories",
    "ifrs-full:CurrentContractAssets": "contract_assets",
    "ifrs-full:CurrentPrepayments": "prepayments",
    "ifrs-full:TradeAndOtherCurrentPayablesToTradeSuppliers":
        "trade_payables",
    "ifrs-full:TradeAndOtherCurrentPayablesToRelatedParties":
        "related_trade_payables",
    "tifrs-bsci-ci:NotesPayable": "notes_payable",
    "ifrs-full:OtherCurrentPayables": "other_current_payables",
    "tifrs-bsci-ci:OtherPayablesToRelatedParties":
        "related_other_payables",
    "ifrs-full:CurrentContractLiabilities": "contract_liabilities",
    "tifrs-bsci-ci:WagesAndSalariesPayable": "wages_payable",
    "tifrs-bsci-ci:EmployeeBonusPayable": "bonus_payable",
    "tifrs-bsci-ci:PayableOnMachineryAndEquipment": "capex_payable",
    "ifrs-full:CurrentDividendPayables": "dividends_payable",
    "ifrs-full:ShorttermBorrowings": "short_term_borrowings",
    "tifrs-bsci-ci:LongtermLiabilitiesCurrentPortion":
        "current_portion_long_term_debt",
    "ifrs-full:LongtermBorrowings": "long_term_borrowings",
    "ifrs-full:NoncurrentPortionOfNoncurrentBondsIssued":
        "noncurrent_bonds",
    "ifrs-full:CurrentLeaseLiabilities": "current_leases",
    "ifrs-full:NoncurrentLeaseLiabilities": "noncurrent_leases",
    "ifrs-full:NoncurrentFinanceLeaseLiabilities": "noncurrent_leases",
    "ifrs-full:NoncontrollingInterests": "noncontrolling_interests",
    "tifrs-bsci-ci:NumberOfSharesInEntityHeldByEntityAndByItsSubsidiaries":
        "treasury_shares",
}
METRICS = (*HISTORY_METRICS, "issued_shares", "extra_reinvestment",
           *dict.fromkeys(TAGS.values()))
METRICS = tuple(dict.fromkeys(METRICS))
SHARE_METRICS = ("shares", "issued_shares", "treasury_shares")
DRIVERS = {
    "revenue_growth": (16, "Revenue growth"),
    "ebit_margin": (21, "EBIT margin"),
    "tax_rate": (26, "Operating tax rate"),
    "da_ratio": (31, "D&A / revenue"),
    "capex_ratio": (36, "Capex / revenue"),
    "nwc_ratio": (41, "Operating NWC / revenue"),
    "extra_reinvestment": (60, "Extra reinvestment / cash use (+/-)"),
}
CAPITAL = {
    "risk_free": (49, "Risk-free rate"),
    "beta": (50, "Levered beta"),
    "equity_risk_premium": (51, "Equity risk premium"),
    "cost_of_debt": (52, "Pre-tax cost of debt"),
    "debt_weight": (53, "Debt / (debt + equity), market values"),
    "wacc_tax_rate": (54, "Marginal tax rate for debt shield"),
    "terminal_growth": (57, "Terminal growth"),
    "terminal_roic": (58, "Terminal return on invested capital"),
}
INPUTS = {
    "base_revenue": (7, "Annual revenue base"),
    "base_nwc": (8, "Opening operating NWC"),
    "cash": (9, "Cash and cash equivalents"),
    "debt": (10, "Interest-bearing debt incl. leases"),
    "shares": (11, "Common shares outstanding (million)"),
    "equity_adjustment": (12, "Other adjustments to common equity (+/-)"),
}
LEGACY_REVIEW_NOTES = ("basis_reviewed", "terminal_reviewed")
NUMBER = '#,##0.0;(#,##0.0);"-"'
RATE = '0.0%;(0.0%);"-"'
BLUE = "1456A0"
GREEN = "217345"
INK = "1F2937"
AMBER = "FFF4D1"


@dataclass(frozen=True)
class Fact:
    """A normalized fact with its source period and original representation."""

    metric: str
    value: float
    start: str
    end: str
    period: str
    source_url: str
    retrieved_at: str
    published_at: str = ""
    output_date: str = ""
    tag: str = ""
    raw_value: str = ""
    raw_unit: str = "TWD"
    scale: int = 0
    sign: str = ""
    scope: str = "consolidated"
    classification: str = "reported"
    method: str = ""


@dataclass
class Dataset:
    """Current company identity, normalized facts and explicit intake gaps."""

    ticker: str
    company: str
    market: str
    industry: str
    valuation_date: date
    facts: list[Fact] = field(default_factory=list)
    gaps: list[str] = field(default_factory=list)
    financial: bool = False
    classification_note: str = ""


def _number(value: Any) -> float:
    """Return a finite number; reject missing values, booleans and text."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError("Expected a finite numeric value")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("Non-finite numeric value")
    return result


def _fetch(url: str, timeout: float = 25) -> bytes:
    """Read one public URL with a bounded timeout and response size.

    Raises HTTPError, URLError or ValueError without substituting data.
    """
    request = Request(url, headers={"User-Agent": "PyInvest-DCF/1.0.0"})
    with urlopen(request, timeout=timeout) as response:
        body = response.read(12_000_001)
    if len(body) > 12_000_000:
        raise ValueError("Public response exceeds 12 MB")
    return body


def _decode(body: bytes) -> str:
    """Decode the official UTF-8 or legacy Big5 HTML without replacement."""
    try:
        return body.decode("utf-8-sig")
    except UnicodeDecodeError:
        return body.decode("cp950")


class _InlineXbrl(HTMLParser):
    """Read inline facts, contexts and units without inferring tag meanings."""

    def __init__(self) -> None:
        """Initialize the parser's bounded-document state."""
        super().__init__(convert_charrefs=True)
        self.contexts: dict[str, dict[str, str | bool]] = {}
        self.units: dict[str, list[str]] = {}
        self.values: list[tuple[dict[str, str], str]] = []
        self.metadata: dict[str, str] = {}
        self.context = ""
        self.unit = ""
        self.capture = ""
        self.parts: list[str] = []
        self.attributes: dict[str, str] = {}

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        """Track XBRL contexts, explicit dimensions, units and inline facts."""
        attributes = {key: value or "" for key, value in attrs}
        if tag == "xbrli:context":
            self.context = attributes.get("id", "")
            self.contexts[self.context] = {"dimensions": False}
        elif tag == "xbrli:unit":
            self.unit = attributes.get("id", "")
            self.units[self.unit] = []
        elif tag in ("xbrldi:explicitmember", "xbrldi:typedmember"):
            if self.context:
                self.contexts[self.context]["dimensions"] = True
        if tag in (
            "xbrli:startdate", "xbrli:enddate", "xbrli:instant",
            "xbrli:identifier", "xbrli:measure", "ix:nonfraction",
            "ix:nonnumeric",
        ):
            self.capture = tag
            self.parts = []
            self.attributes = attributes

    def handle_data(self, data: str) -> None:
        """Collect content only for the currently relevant XBRL tag."""
        if self.capture:
            self.parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        """Finish contexts, units and values while preserving raw content."""
        if tag == self.capture:
            value = "".join(self.parts).strip()
            if tag in ("ix:nonfraction", "ix:nonnumeric"):
                if tag == "ix:nonfraction":
                    self.values.append((self.attributes, value))
                else:
                    self.metadata[self.attributes.get("name", "")] = value
            elif tag == "xbrli:measure" and self.unit:
                self.units[self.unit].append(value)
            elif self.context:
                self.contexts[self.context][tag.split(":")[1]] = value
            self.capture = ""
        if tag == "xbrli:context":
            self.context = ""
        elif tag == "xbrli:unit":
            self.unit = ""


def inspect_ixbrl(
    body: bytes, ticker: str, year: int
) -> list[dict[str, Any]]:
    """Inspect source concepts and contexts without assigning account meaning.

    Return original tags, periods, entity, dimensions, units and raw numbers.
    Only TWD and share units with known transforms get normalized values.
    Unmapped concepts stay available for an agent's filing/note research.
    """
    parser = _InlineXbrl()
    parser.feed(_decode(body))
    if parser.metadata.get("tifrs-notes:CompanyID") != ticker:
        raise ValueError("XBRL company identifier does not match request")
    category = parser.metadata.get("tifrs-notes:ReportCategory", "")
    rows: list[dict[str, Any]] = []
    for attributes, raw in parser.values:
        tag = attributes.get("name", "")
        context = parser.contexts.get(attributes.get("contextref", ""), {})
        unit = parser.units.get(attributes.get("unitref", ""), [])
        start = str(context.get("startdate", ""))
        end = str(context.get("instant") or context.get("enddate", ""))
        if not end.startswith(f"{year}-"):
            continue
        date.fromisoformat(end)
        if start:
            date.fromisoformat(start)
        if not start:
            period = "instant"
        elif start == f"{year}-01-01":
            period = "annual" if end == f"{year}-12-31" else "ytd"
        else:
            period = "quarter"
        value = None
        scale = 0
        try:
            scale = int(attributes.get("scale", "0"))
            if (unit in (["iso4217:TWD"], ["xbrli:shares"])
                    and attributes.get("format", "") in (
                        "ixt:numdotdecimal", "ixt:numcommadot", "")
                    and attributes.get("xsi:nil") != "true"):
                numeric = float(raw.replace(",", "").replace(" ", ""))
                numeric *= (-1 if attributes.get("sign") == "-" else 1)
                value = _number(numeric * 10 ** scale / 1_000_000)
        except (ValueError, OverflowError):
            value = None
        rows.append({
            "tag": tag, "context": attributes.get("contextref", ""),
            "entity": context.get("identifier", ""),
            "dimensions": bool(context.get("dimensions")),
            "start": start, "end": end, "period": period,
            "raw_value": raw, "unit": unit, "scale": scale,
            "sign": attributes.get("sign", ""), "value": value,
            "report_category": category,
        })
    return rows


def parse_ixbrl(
    body: bytes, ticker: str, year: int, url: str, retrieved_at: str
) -> list[Fact]:
    """Map verified consolidated TWD/share concepts; retain absent components.

    Components are not blindly summed into NWC, debt or equity. Dimensions,
    foreign currencies and conflicting duplicates do not enter the model.
    """
    candidates: dict[tuple[str, str, str], Fact] = {}
    conflicts: set[tuple[str, str, str]] = set()
    for row in inspect_ixbrl(body, ticker, year):
        category = row["report_category"]
        if "consolidated" not in category.lower() and "合併" not in category:
            raise ValueError("Consolidated report scope is not confirmed")
        metric = TAGS.get(row["tag"])
        expected = (["xbrli:shares"] if metric in SHARE_METRICS else
                    ["iso4217:TWD"])
        if (not metric or row["dimensions"] or row["entity"] != ticker
                or row["unit"] != expected or row["value"] is None):
            continue
        value = row["value"]
        if metric.startswith("capex_"):
            value = abs(value)
        key = (metric, row["start"], row["end"])
        fact = Fact(
            metric, value, row["start"], row["end"], row["period"],
            url, retrieved_at, tag=row["tag"], raw_value=row["raw_value"],
            scale=row["scale"], sign=row["sign"],
            raw_unit="shares" if metric in SHARE_METRICS else "TWD",
        )
        if key in candidates and candidates[key].value != value:
            conflicts.add(key)
        candidates[key] = fact
    return [fact for key, fact in candidates.items() if key not in conflicts]


def annual_value(facts: list[Fact], metric: str, year: int) -> float | None:
    """Return annual flow or year-end stock without summing YTD overlaps.

    Prefer an explicit annual fact. Otherwise require four exact quarters,
    converting YTD to single quarters only with the prior YTD available.
    Conflicting duplicate periods raise ValueError; missing remains None.
    """
    relevant = [
        fact for fact in facts if fact.metric == metric
        and fact.end.startswith(f"{year}-")
    ]
    keyed: dict[tuple[str, str, str], float] = {}
    for fact in relevant:
        key = (fact.period, fact.start, fact.end)
        if key in keyed and keyed[key] != fact.value:
            raise ValueError(f"Conflicting {metric} facts for {key}")
        keyed[key] = fact.value
    for kind in ("annual", "instant"):
        key = (kind, f"{year}-01-01" if kind == "annual" else "",
               f"{year}-12-31")
        if key in keyed:
            return keyed[key]
    quarters: list[float] = []
    previous_ytd: float | None = 0.0
    for quarter, end in enumerate(("03-31", "06-30", "09-30", "12-31"), 1):
        start = f"{year}-{(quarter - 1) * 3 + 1:02d}-01"
        single = keyed.get(("quarter", start, f"{year}-{end}"))
        ytd = keyed.get(("ytd", f"{year}-01-01", f"{year}-{end}"))
        if single is None and ytd is not None and previous_ytd is not None:
            single = ytd - previous_ytd
        if single is None:
            return None
        quarters.append(single)
        if ytd is not None and not math.isclose(
                sum(quarters), ytd, rel_tol=1e-10, abs_tol=1e-8):
            raise ValueError("Quarter and YTD facts disagree")
        previous_ytd = sum(quarters)
    return sum(quarters) if len(quarters) == 4 else None


def ttm_value(facts: list[Fact], metric: str, end: str) -> float | None:
    """Derive TTM flow = prior FY + current YTD - matching prior YTD.

    Require the same calendar cutoff, scope and normalized units. Missing
    comparatives remain unavailable; this never annualizes a short period.
    """
    if metric not in HISTORY_METRICS[:8]:
        raise ValueError("TTM is only defined for the supported flow metrics")
    _validate_facts([fact for fact in facts if fact.metric == metric])
    cutoff = date.fromisoformat(end)
    if cutoff.month == 12 and cutoff.day == 31:
        return annual_value(facts, metric, cutoff.year)
    if (cutoff.month, cutoff.day) not in ((3, 31), (6, 30), (9, 30)):
        raise ValueError("TTM requires an exact calendar quarter end")
    annual = annual_value(facts, metric, cutoff.year - 1)
    values: list[float] = []
    for year in (cutoff.year, cutoff.year - 1):
        matches = {fact.value for fact in facts if fact.metric == metric
                   and fact.start == f"{year}-01-01"
                   and fact.end == f"{year}-{end[5:]}"
                   and fact.period in ("ytd", "quarter")}
        if len(matches) > 1:
            raise ValueError("Conflicting TTM comparatives")
        if not matches:
            return None
        values.append(matches.pop())
    return annual + values[0] - values[1] if annual is not None else None


def _roster(market: str) -> tuple[list[dict[str, str]], str]:
    """Read the official company roster, with official CSV fallback."""
    url = (TWSE + "t187ap03_L" if market == "TWSE" else
           TPEX + "mopsfin_t187ap03_O")
    try:
        rows = json.loads(_decode(_fetch(url)))
        if not isinstance(rows, list):
            raise ValueError("Roster response is not a list")
    except (HTTPError, URLError, ValueError, UnicodeError):
        suffix = "L" if market == "TWSE" else "O"
        url = f"https://mopsfin.twse.com.tw/opendata/t187ap03_{suffix}.csv"
        rows = list(csv.DictReader(io.StringIO(_decode(_fetch(url)))))
    return rows, url


def collect_current_data(
    ticker: str, capture_dir: Path | None = None
) -> Dataset:
    """Fetch five completed calendar years plus the latest disclosed quarter.

    Only public official endpoints are used. Individual filing failures are
    recorded as gaps. A missing identity or inaccessible rosters raises
    ValueError/URLError. Capture files are optional local reproducibility data.
    """
    if not re.fullmatch(r"\d{4,6}", ticker):
        raise ValueError("Use a numeric TWSE or TPEx company code")
    today = date.today()
    retrieved = datetime.now(timezone.utc).isoformat()
    identity: dict[str, str] | None = None
    errors: list[str] = []
    roster_url = ""
    market = ""
    for market in ("TWSE", "TPEx"):
        try:
            rows, roster_url = _roster(market)
        except (HTTPError, URLError, ValueError, UnicodeError) as error:
            errors.append(f"{market} roster: {type(error).__name__}: {error}")
            continue
        identity = next((row for row in rows if
                         row.get("公司代號", row.get(
                             "SecuritiesCompanyCode")) == ticker), None)
        if identity:
            break
    if not identity:
        detail = "; ".join(errors) or "Code absent from both current rosters"
        raise ValueError(detail)
    industry = identity.get("產業別", identity.get("SecuritiesIndustryCode", ""))
    company = identity.get("公司名稱", identity.get("CompanyName", ticker))
    data = Dataset(ticker, company, market, industry, today, gaps=errors)
    data.financial = industry == "17"
    data.classification_note = (
        f"Official roster industry code: {industry}; {roster_url}"
        if industry else "Industry missing; provisional nonfinancial "
        "classification requires business-description research.")
    if not industry:
        data.gaps.append("Company industry classification is missing")
    if capture_dir:
        capture_dir.mkdir(parents=True, exist_ok=True)
        identity_keys = (
            "公司代號", "公司名稱", "產業別", "出表日期",
            "已發行普通股數或TDR原股發行股數", "SecuritiesCompanyCode",
            "CompanyName", "SecuritiesIndustryCode", "Date", "IssueShares",
        )
        safe_identity = {key: value for key, value in identity.items()
                         if key in identity_keys}
        (capture_dir / f"{ticker}-identity.json").write_text(json.dumps(
            {"url": roster_url, "retrieved_at": retrieved,
             "identity": safe_identity}, ensure_ascii=False, indent=2),
            encoding="utf-8")
    if data.financial:
        data.gaps.append("Financial industry: general FCFF is not applicable")
        return data
    raw_shares = identity.get(
        "已發行普通股數或TDR原股發行股數", identity.get("IssueShares"))
    if raw_shares is not None:
        try:
            issued = _number(float(str(raw_shares).replace(",", "")))
            if issued < 0:
                raise ValueError("Negative issued shares")
            data.facts.append(Fact(
                "issued_shares", issued / 1_000_000, "", today.isoformat(),
                "snapshot", roster_url, retrieved,
                output_date=str(identity.get(
                    "出表日期", identity.get("Date", ""))),
                raw_value=str(raw_shares), raw_unit="shares",
                method="Roster issued shares observed on snapshot date; "
                "not automatically outstanding common shares. Check TDR "
                "basis, treasury shares and later changes.",
            ))
        except (ValueError, TypeError):
            data.gaps.append("Roster issued-share value could not be parsed")
    filings = [(year, 4) for year in range(today.year - 5, today.year)]
    income_url = (TWSE + "t187ap06_L_ci" if market == "TWSE" else
                  TPEX + "mopsfin_t187ap06_O_ci")
    try:
        income = json.loads(_decode(_fetch(income_url)))
        row = next((item for item in income if str(
            item.get("公司代號") or item.get("CompanyCode")
            or item.get("SecuritiesCompanyCode", "")) == ticker), None)
        if row:
            raw_year = int(row.get("年度", row.get("Year", "0")))
            year = raw_year + 1911 if raw_year < 1911 else raw_year
            quarter = int(row.get("季別", row.get(
                "Season", row.get("Quarter", "0"))))
            if year >= today.year - 1 and 1 <= quarter <= 4:
                if (year, quarter) not in filings:
                    filings.append((year, quarter))
            if capture_dir:
                (capture_dir / f"{ticker}-latest-snapshot.json").write_text(
                    json.dumps({"url": income_url, "retrieved_at": retrieved,
                                "row": row}, ensure_ascii=False, indent=2),
                    encoding="utf-8")
    except (HTTPError, URLError, ValueError, UnicodeError) as error:
        data.gaps.append(f"Latest quarter: {type(error).__name__}: {error}")
    completed_quarter = (today.month - 1) // 3
    if completed_quarter and (today.year, completed_quarter) not in filings:
        # One relevant probe; a missing filing still requires issuer IR review.
        filings.append((today.year, completed_quarter))
    for year, quarter in filings:
        url = MOPS + "?" + urlencode({
            "step": 1, "CO_ID": ticker, "SYEAR": year,
            "SSEASON": quarter, "REPORT_ID": "C",
        })
        try:
            body = _fetch(url)
            if capture_dir:
                (capture_dir / f"{ticker}-{year}Q{quarter}.html").write_bytes(
                    body)
            parsed = parse_ixbrl(body, ticker, year, url, retrieved)
            if not parsed:
                raise ValueError("No supported consolidated TWD facts")
            data.facts.extend(parsed)
            if quarter < 4:
                data.facts.extend(fact for fact in parse_ixbrl(
                    body, ticker, year - 1, url, retrieved)
                    if fact.start and not fact.end.endswith("12-31"))
        except (HTTPError, URLError, ValueError, UnicodeError) as error:
            data.gaps.append(
                f"{year}Q{quarter}: {type(error).__name__}: {error}")
    data.gaps.extend([
        "Operating NWC needs a company-specific operating-account definition",
        "Debt must include all interest-bearing debt and relevant leases",
        "Common-equity adjustments and outstanding shares require review",
        "Agent must research or explicitly estimate capital assumptions",
        "Announcement dates are unknown unless supplied separately",
        "MOPS may lag issuer IR: agent must check latest published earnings",
    ])
    if capture_dir:
        (capture_dir / "normalized.json").write_text(json.dumps(
            {**asdict(data), "valuation_date": today.isoformat()},
            ensure_ascii=False, indent=2), encoding="utf-8")
    return data


def load_supplement(path: Path) -> dict[str, Any]:
    """Read documented JSON facts and evidenced inputs; reject invalid values.

    Supplemental facts already use TWD million or million shares. They do not
    silently override conflicting official facts. Raises ValueError on schema
    or finite-number violations and OSError on unreadable files.
    """
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Supplement must be a JSON object")
    for row in data.get("facts", []):
        fact = Fact(**row)
        _number(fact.value)
        if fact.metric not in METRICS or fact.period not in (
            "annual", "quarter", "ytd", "instant", "snapshot",
        ):
            raise ValueError("Unknown supplemental metric or period")
        date.fromisoformat(fact.end)
        if fact.start:
            date.fromisoformat(fact.start)
        if not fact.source_url or not fact.retrieved_at:
            raise ValueError("Fact source and retrieval time are required")
        if fact.scope != "consolidated":
            raise ValueError("Supplement must use consolidated scope")
    for key, item in data.get("inputs", {}).items():
        if key not in INPUTS and key not in LEGACY_REVIEW_NOTES:
            raise ValueError(f"Unknown input: {key}")
        _evidenced(item)
        if isinstance(item["value"], list):
            raise ValueError("Model inputs must be scalar values")
    for case, values in data.get("scenarios", {}).items():
        if case not in CASES:
            raise ValueError(f"Unknown case: {case}")
        for key, item in values.items():
            if key not in DRIVERS and key not in CAPITAL:
                raise ValueError(f"Unknown driver: {key}")
            value = _evidenced(item)
            if key in DRIVERS and (not isinstance(value, list)
                                   or len(value) != 5):
                raise ValueError("Forecast drivers need five annual values")
            if key in CAPITAL and isinstance(value, list):
                raise ValueError("Capital assumptions must be scalars")
    for entry in data.get("research_log", []):
        if (not isinstance(entry, dict) or not entry.get("topic")
                or not entry.get("source") or not entry.get("result")
                or not entry.get("notes")):
            raise ValueError("Research log needs topic/source/result/notes")
    return data


def _evidenced(item: dict[str, Any]) -> float | list[float]:
    """Validate classification, provenance, dates and optional components."""
    if not item.get("source") or not item.get("method"):
        raise ValueError("Every assumption needs source and method")
    item.setdefault("classification", "analyst_estimate")
    if item["classification"] not in CLASSIFICATIONS:
        raise ValueError("Assumption classification is required")
    for key in ("as_of", "period_start", "period_end"):
        if item.get(key):
            date.fromisoformat(item[key])
    value = item.get("value")
    components = item.get("components", [])
    if components:
        if isinstance(value, list):
            raise ValueError("Components only reconcile scalar model inputs")
        total = 0.0
        for part in components:
            if (not part.get("label") or isinstance(part.get("sign"), bool)
                    or part.get("sign") not in (-1, 1)):
                raise ValueError("Components need label and explicit +/- sign")
            total += _number(_evidenced(part)) * part["sign"]
        if not math.isclose(total, _number(value), abs_tol=1e-6, rel_tol=1e-9):
            raise ValueError("Component sum does not match supplied total")
    if isinstance(value, list):
        return [_number(number) for number in value]
    return _number(value)


def _validate_facts(facts: list[Fact]) -> None:
    """Reject invalid periods and conflicts across every source period."""
    values: dict[tuple[str, str, str, str], float] = {}
    for fact in facts:
        _number(fact.value)
        if (fact.metric not in METRICS or fact.scope != "consolidated"
                or fact.classification not in CLASSIFICATIONS):
            raise ValueError("Unknown metric or non-consolidated fact")
        end = date.fromisoformat(fact.end)
        start = date.fromisoformat(fact.start) if fact.start else None
        if fact.period in ("instant", "snapshot"):
            valid = start is None
        elif fact.period == "annual":
            valid = start == date(end.year, 1, 1) and end == date(
                end.year, 12, 31)
        elif fact.period == "ytd":
            valid = start == date(end.year, 1, 1)
        elif fact.period == "quarter":
            valid = (start is not None and (end.month, end.day) in (
                (3, 31), (6, 30), (9, 30), (12, 31)) and start == date(
                    end.year, end.month - 2, 1))
        else:
            valid = False
        if not valid:
            raise ValueError("Fact period does not match its dates")
        key = (fact.metric, fact.period, fact.start, fact.end)
        if key in values and values[key] != fact.value:
            raise ValueError(f"Conflicting source values for {key}")
        values[key] = fact.value


def _cell(sheet: Worksheet, row: int, col: int, value: Any) -> None:
    """Write typed cells while preventing untrusted text formula injection."""
    cell = sheet.cell(row, col)
    cell.value = value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        cell.data_type = "s"


def _formula(sheet: Worksheet, row: int, col: int, value: str) -> None:
    """Write a trusted formula and use finance formula colors."""
    cell = sheet.cell(row, col, value)
    cell.font = Font(name="Arial", size=10, color=(
        GREEN if "!" in value else INK))


def _input(sheet: Worksheet, row: int, col: int, value: Any = None) -> None:
    """Create an editable input with a clear missing-value fill."""
    _cell(sheet, row, col, value)
    cell = sheet.cell(row, col)
    cell.fill = PatternFill("solid", fgColor=AMBER)
    cell.font = Font(name="Arial", size=10, color=BLUE)


def _select(sheet: Worksheet, row: int, col: int, range_: str) -> None:
    """Select one scenario without converting its blank input into zero."""
    _formula(sheet, row, col, (
        '=IF(ISNUMBER($C$4),IF(AND($C$4=INT($C$4),$C$4>=1,$C$4<=3),'
        f'IF(COUNT(INDEX({range_},$C$4))=1,INDEX({range_},$C$4),""),""),"")'
    ))


def _style(sheet: Worksheet, last_col: int = 9) -> None:
    """Apply compact, readable worksheet defaults to a bounded range."""
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C4"
    sheet.column_dimensions["A"].width = 2
    sheet.column_dimensions["B"].width = 48
    for column in range(3, last_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.column_dimensions["H"].width = 3
    sheet.column_dimensions["I"].width = 78
    for row in sheet.iter_rows():
        sheet.row_dimensions[row[0].row].height = 19
        for cell in row:
            if not cell.has_style:
                cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="center")
            if cell.column > 2:
                cell.number_format = NUMBER
    sheet["B2"].font = Font(name="Arial", size=14, bold=True, color=INK)


def _header(sheet: Worksheet, row: int, labels: list[str | int]) -> None:
    """Write a bounded column header with high contrast."""
    for col, label in enumerate(labels, 2):
        cell = sheet.cell(row, col, label)
        cell.fill = PatternFill("solid", fgColor="203548")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if isinstance(label, int):
            cell.number_format = "0"


def _history(book: Workbook, data: Dataset, facts: list[Fact]) -> None:
    """Lay out five annual periods and source-linked calibration formulas."""
    sheet = book["History"]
    sheet["B2"] = "Five-year consolidated history (TWD million)"
    years = range(data.valuation_date.year - 5, data.valuation_date.year)
    _header(sheet, 3, ["Metric", *years])
    for row, metric in enumerate(HISTORY_METRICS, 4):
        sheet.cell(row, 2, metric)
        for col, year in enumerate(years, 3):
            value = annual_value(facts, metric, year)
            _input(sheet, row, col, value)
    for row, label in ((18, "D&A"), (19, "Cash capex incl. intangibles"),
                       (20, "EBIT margin"), (21, "Effective tax rate"),
                       (22, "Historical FCFF proxy (effective tax)"),
                       (23, "Opening operating NWC")):
        sheet.cell(row, 2, label)
    for col in range(3, 8):
        letter = get_column_letter(col)
        _input(sheet, 24, col, annual_value(
            facts, "extra_reinvestment",
            data.valuation_date.year - 5 + col - 3))
        for row, cells in ((18, (8, 9)), (19, (10, 11))):
            refs = ",".join(f"{letter}{number}" for number in cells)
            _formula(sheet, row, col,
                     f'=IF(COUNT({refs})=2,SUM({refs}),"")')
        _formula(sheet, 20, col, (
            f'=IF(COUNT({letter}4:{letter}5)=2,'
            f'IF({letter}4>0,{letter}5/{letter}4,""),"")'))
        _formula(sheet, 21, col, (
            f'=IF(COUNT({letter}6:{letter}7)=2,'
            f'IF({letter}6>0,{letter}7/{letter}6,""),"")'))
        sheet.cell(20, col).number_format = RATE
        sheet.cell(21, col).number_format = RATE
        if col == 3:
            opening = annual_value(
                facts, "nwc", data.valuation_date.year - 6)
            _input(sheet, 23, col, opening)
        else:
            previous = get_column_letter(col - 1)
            _formula(sheet, 23, col, (
                f'=IF(ISNUMBER({previous}12),{previous}12,"")'))
        _formula(sheet, 22, col, (
            f'=IF(COUNT({letter}5,{letter}12,{letter}18:{letter}19,'
            f'{letter}21,{letter}23:{letter}24)=7,'
            f'IF(AND({letter}21>=0,{letter}21<=1,'
            f'{letter}18>=0,{letter}19>=0),'
            f'{letter}5-MAX({letter}5,0)*{letter}21+'
            f'{letter}18-{letter}19-({letter}12-{letter}23)'
            f'-{letter}24,""),"")'))
    sheet["I4"] = "Only full-year flows and year-end balances appear here."
    sheet["I5"] = "Raw periods, tags, units and provenance are in Facts."
    sheet["I6"] = "Annual flows take precedence over overlapping YTD."
    sheet["I7"] = "Manual edits require a corresponding source in Facts."
    sheet["I8"] = "NWC is operating assets less operating liabilities."
    sheet["I9"] = "Total liabilities are not interest-bearing debt."
    sheet["I22"] = "Tax proxy = tax expense / pretax profit; no loss refund."
    sheet["I23"] = "FCFF needs prior-year NWC; a missing balance stays blank."
    sheet["B24"] = "Extra operating reinvestment / cash use (+/-)"
    sheet["I24"] = (
        "BOT / construction / prepaid cash may exceed PPE purchases.")
    sheet.sheet_properties.tabColor = "E1D1BD"


def _facts(book: Workbook, facts: list[Fact], gaps: list[str]) -> None:
    """Preserve source facts, intake gaps and unknown metadata."""
    sheet = book["Facts"]
    headers = list(Fact.__dataclass_fields__)
    for col, name in enumerate(headers, 1):
        sheet.cell(1, col, name)
        sheet.column_dimensions[get_column_letter(col)].width = (
            65 if name in ("source_url", "tag") else 24)
    for row, fact in enumerate(facts, 2):
        for col, value in enumerate(asdict(fact).values(), 1):
            _cell(sheet, row, col, value)
    gap_start = len(facts) + 4
    sheet.cell(
        gap_start, 1, "Initial intake observations; see research results")
    for row, gap in enumerate(gaps, gap_start + 1):
        _cell(sheet, row, 1, gap)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(1, len(facts) + 1)}")
    sheet.sheet_view.showGridLines = False


def _source_note(item: dict[str, Any]) -> str:
    """Keep classification, dates, source and method visible together."""
    note = (f'{item["classification"]}: {item["source"]}; '
            f'{item["method"]}')
    for key in ("as_of", "period_start", "period_end"):
        if item.get(key):
            note += f"; {key}={item[key]}"
    return note


def _research(book: Workbook, supplement: dict[str, Any]) -> None:
    """Show component reconciliations and the agent's source-search record."""
    sheet = book["Evidence"]
    sheet["B2"] = "Input components and research record"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C5"
    for column, width in (("A", 2), ("B", 42), ("C", 18), ("D", 16),
                          ("E", 22), ("F", 18), ("G", 64), ("H", 88),
                          ("I", 18), ("J", 18)):
        sheet.column_dimensions[column].width = width
    row = 4
    for key, item in supplement.get("inputs", {}).items():
        components = item.get("components", [])
        if not components:
            continue
        input_row = INPUTS[key][0]
        _cell(sheet, row, 2, key)
        _formula(sheet, row, 3, f"=Assumptions!C{input_row}")
        sheet.cell(row, 4, "Input total")
        sheet.cell(row, 5, "Component total")
        sheet.cell(row, 7, "Difference (component - input)")
        first = row + 2
        last = first + len(components) - 1
        _formula(sheet, row, 6,
                 f"=SUMPRODUCT(C{first}:C{last},D{first}:D{last})")
        _formula(sheet, row, 8, f"=F{row}-C{row}")
        sheet.conditional_formatting.add(
            f"H{row}", CellIsRule(operator="notBetween",
                                  formula=[-0.000001, 0.000001],
                                  fill=PatternFill("solid", fgColor="FCE7E7")))
        _header(sheet, row + 1, [
            "Component", "Value", "Sign", "Classification", "As of",
            "Source", "Method", "Period start", "Period end",
        ])
        for offset, part in enumerate(components, first):
            values = [part["label"], part["value"], part["sign"],
                      part["classification"], part.get("as_of", ""),
                      part["source"], part["method"],
                      part.get("period_start", ""), part.get("period_end", "")]
            for col, value in enumerate(values, 2):
                _cell(sheet, offset, col, value)
            sheet.row_dimensions[offset].height = 48
        row = last + 3
    sheet.cell(row, 2, "Research log: inspect / derive / estimate / missing")
    _header(sheet, row + 1,
            ["Topic", "Result", "Retrieved at", "Source", "Notes"])
    for index, entry in enumerate(supplement.get("research_log", []), row + 2):
        for col, key in enumerate(
                ("topic", "result", "retrieved_at", "source", "notes"), 2):
            _cell(sheet, index, col, entry.get(key, ""))
        sheet.row_dimensions[index].height = 48
    for cells in sheet.iter_rows():
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = NUMBER


def _assumptions(
    book: Workbook, data: Dataset, facts: list[Fact],
    supplement: dict[str, Any],
) -> None:
    """Create three editable cases and one authoritative active selection."""
    sheet = book["Assumptions"]
    sheet["B2"] = "Assumptions (TWD million; million shares)"
    sheet["B4"] = "Case: 1 Bear / 2 Base / 3 Bull"
    _input(sheet, 4, 3, 2)
    sheet["C4"].number_format = "0"
    _formula(sheet, 4, 4,
             '=IF(ISNUMBER(C4),IF(AND(C4=INT(C4),C4>=1,C4<=3),'
             'CHOOSE(C4,"Bear","Base","Bull"),"Invalid case"),'
             '"Invalid case")')
    selector = DataValidation(type="whole", operator="between",
                              formula1=1, formula2=3, allow_blank=False)
    selector.error = "Choose 1, 2 or 3"
    selector.showErrorMessage = True
    sheet.add_data_validation(selector)
    selector.add("C4")
    sheet["B6"] = "Evidence grade"
    for key, (row, label) in INPUTS.items():
        sheet.cell(row, 2, label)
        item = supplement.get("inputs", {}).get(key)
        _input(sheet, row, 3, item["value"] if item else None)
        if item:
            _cell(sheet, row, 9, _source_note(item))
    if "base_revenue" not in supplement.get("inputs", {}):
        _formula(sheet, 7, 3, '=IF(ISNUMBER(History!G4),History!G4,"")')
        sheet["I7"] = (
            "analyst_estimate: History latest FY; draft revenue base, "
            "review for current use and search latest IR / TTM.")
    current_cash = sorted((fact for fact in facts if fact.metric == "cash"
                           and fact.period == "instant"),
                          key=lambda fact: fact.end, reverse=True)
    if current_cash and "cash" not in supplement.get("inputs", {}):
        cash = current_cash[0]
        _input(sheet, 9, 3, cash.value)
        note = (f"{cash.classification}: Cash at {cash.end}; "
                f"{cash.source_url}")
        if cash.method:
            note += f"; {cash.method}"
        _cell(sheet, 9, 9, note)
    if "shares" not in supplement.get("inputs", {}):
        sheet["I11"] = "Outstanding common shares; not EPS weighted shares."
    if "equity_adjustment" not in supplement.get("inputs", {}):
        sheet["I12"] = "Include NCI, preferred equity, non-operating assets."
    base_date = supplement.get("inputs", {}).get("base_revenue", {}).get(
        "period_end", "not supplied")
    nwc_date = supplement.get("inputs", {}).get("base_nwc", {}).get(
        "as_of", "not supplied")
    sheet["I13"] = f"Revenue cutoff {base_date}; opening NWC {nwc_date}."
    sheet["I14"] = (
        "Explain terminal ROIC and reinvestment transition; no checkbox.")
    for key, (row, label) in DRIVERS.items():
        sheet.cell(row, 2, f"{label} / Active")
        for case_index, case in enumerate(CASES, 1):
            sheet.cell(row + case_index, 2, case.title())
            item = supplement.get("scenarios", {}).get(case, {}).get(key)
            for col in range(3, 8):
                _input(sheet, row + case_index, col,
                       item["value"][col - 3] if item else None)
            if item:
                _cell(sheet, row + case_index, 9, _source_note(item))
            elif case == "base" and key in (
                    "ebit_margin", "tax_rate", "da_ratio", "capex_ratio"):
                # Flat actual-based drafts are visible and freely editable.
                history_row = {"ebit_margin": 20, "tax_rate": 21,
                               "da_ratio": 18, "capex_ratio": 19}[key]
                ratio = key in ("da_ratio", "capex_ratio")
                expression = f"History!G{history_row}"
                test = f"ISNUMBER({expression})"
                if ratio:
                    test = f"AND({test},ISNUMBER(History!G4),History!G4>0)"
                    expression += "/History!G4"
                for col in range(3, 8):
                    _formula(sheet, row + case_index, col,
                             f'=IF({test},{expression},"")')
                _cell(sheet, row + case_index, 9, (
                    f"analyst_estimate: DRAFT latest FY History!G{history_row}"
                    f"{' / History!G4' if ratio else ''}; flat for 5 years; "
                    "review before use. Sources in Facts."))
        for col in range(3, 8):
            letter = get_column_letter(col)
            _select(sheet, row, col, f"{letter}{row+1}:{letter}{row+3}")
            for number in range(row, row + 4):
                sheet.cell(number, col).number_format = (
                    NUMBER if key == "extra_reinvestment" else RATE)
        _formula(sheet, row, 9, (
            '=IF(ISNUMBER($C$4),IF(AND($C$4=INT($C$4),$C$4>=1,$C$4<=3),'
            f'IF(INDEX(I{row+1}:I{row+3},$C$4)="","",'
            f'INDEX(I{row+1}:I{row+3},$C$4)),""),"")'))
    _header(sheet, 15, ["Operating drivers", "Year 1", "Year 2", "Year 3",
                        "Year 4", "Year 5"])
    _header(sheet, 48, ["Capital assumptions",
            "Active", "Bear", "Base", "Bull"])
    for key, (row, label) in CAPITAL.items():
        sheet.cell(row, 2, label)
        for col, case in enumerate(CASES, 4):
            item = supplement.get("scenarios", {}).get(case, {}).get(key)
            _input(sheet, row, col, item["value"] if item else None)
            if item:
                _cell(sheet, row, 9 + col - 4, _source_note(item))
        # INDEX over a horizontal reference uses an explicit column argument.
        _formula(sheet, row, 3, (
            '=IF(ISNUMBER($C$4),IF(AND($C$4=INT($C$4),$C$4>=1,$C$4<=3),'
            f'IF(COUNT(INDEX(D{row}:F{row},1,$C$4))=1,'
            f'INDEX(D{row}:F{row},1,$C$4),""),""),"")'))
        for col in range(3, 7):
            sheet.cell(row, col).number_format = RATE if row != 50 else NUMBER
        _formula(sheet, row, 8, (
            '=IF(ISNUMBER($C$4),IF(AND($C$4=INT($C$4),$C$4>=1,$C$4<=3),'
            f'IFERROR(LEFT(INDEX(I{row}:K{row},1,$C$4),'
            f'FIND(": ",INDEX(I{row}:K{row},1,$C$4))-1),""),""),"")'))
    sheet["B55"] = "Cost of equity"
    _formula(sheet, 55, 3,
             '=IF(COUNT(C49:C51)=3,C49+C50*C51,"")')
    sheet["B56"] = "WACC"
    _formula(sheet, 56, 3, (
        '=IF(COUNT(C49:C54)=6,IF(AND(C49>=0,C51>=0,C52>=0,'
        'C53>=0,C53<=1,C54>=0,C54<=1),'
        'C55*(1-C53)+C52*(1-C54)*C53,""),"")'))
    sheet["C55"].number_format = RATE
    sheet["C56"].number_format = RATE
    sheet["I56"] = "WACC = (rf + beta * ERP) * (1-wd) + kd * (1-tax) * wd"
    sheet["I65"] = "Terminal growth / ROIC use disclosed proxies or estimates."
    sheet["I66"] = "Extra cash use is after tax; avoid capex / NWC overlap."
    source_refs = ["$I$7:$I$12", *(f"$I${row}" for row, _ in DRIVERS.values())]
    class_refs = ["$H$49:$H$54", "$H$57:$H$58"]
    counts: dict[str, str] = {}
    for classification in CLASSIFICATIONS:
        terms = [f'COUNTIF({ref},"{classification}: *")'
                 for ref in source_refs]
        terms += [f'COUNTIF({ref},"{classification}")' for ref in class_refs]
        counts[classification] = "SUM(" + ",".join(terms) + ")"
    valid_count = "+".join(counts.values())
    expected = len(INPUTS) + len(DRIVERS) + len(CAPITAL)
    _formula(sheet, 6, 3, (
        f'=IF({valid_count}<>{expected},"Rough estimate; partial metadata",'
        f'IF({counts["analyst_estimate"]}>0,"Screen-grade analyst estimate",'
        '"Reported / derived inputs"))'))
    sheet.column_dimensions["H"].width = 22
    sheet["H48"] = "Active classification"
    sheet.sheet_properties.tabColor = "49708A"


def _model(book: Workbook, data: Dataset) -> None:
    """Build one rolling five-year FCFF schedule and formula sensitivities."""
    sheet = book["Model"]
    _cell(sheet, 2, 2, f"{data.ticker} {data.company} DCF")
    sheet["B4"] = "Selected case"
    _formula(sheet, 4, 3, "=Assumptions!D4")
    sheet["B6"] = "Common equity per share (TWD, rough)"
    _formula(sheet, 6, 3, '=IF(ISNUMBER(C42),C42,"")')
    sheet["C6"].number_format = '0.00;(0.00);"-"'
    sheet["B7"] = "Valuation status"
    _formula(sheet, 7, 3, (
        '=IF(C16<>1,"Financial company: general FCFF not applicable",'
        'IF(COUNT(Assumptions!C7:C12)<>6,"Missing base / equity bridge",'
        'IF(COUNT(C29:G29)<>5,"Missing / invalid forecast drivers",'
        'IF(NOT(ISNUMBER(C42)),"Missing / invalid WACC, g, ROIC or bridge",'
        'IF(AND(ISNUMBER(C36),C36<=0),'
        '"Diagnostic estimate: non-positive terminal FCFF",'
        'IF(AND(ISNUMBER(C41),C41<0),'
        '"Diagnostic estimate: negative common equity",'
        'Assumptions!C6))))))'))
    sheet["B9"] = "Valuation date"
    sheet["C9"] = data.valuation_date
    sheet["C9"].number_format = "yyyy-mm-dd"
    sheet["B10"] = "Currency / scale"
    sheet["C10"] = "TWD million"
    sheet["B11"] = "Forecast timing"
    sheet["C11"] = "Rolling years"
    sheet["B12"] = "Evidence grade"
    _formula(sheet, 12, 3, "=Assumptions!C6")
    sheet["I9"] = "Year-end discounting at annual anniversaries of valuation."
    sheet["I10"] = "Revenue base is a reviewed annual run-rate, not YTD."
    sheet["I11"] = "Tax = MAX(EBIT, 0) * tax rate; no loss carryforward."
    sheet["I12"] = "Terminal FCFF = NOPAT6 * (1-g/ROIC); see transition below."
    sheet["I13"] = "Changing Assumptions!C4 recalculates this single forecast."
    sheet["I14"] = "Blank inputs are unavailable; an explicit zero is valid."
    sheet["B16"] = "General FCFF scope (research classification)"
    sheet["C16"] = 0 if data.financial else 1
    sheet["C16"].number_format = "0"
    _cell(sheet, 16, 9, data.classification_note or (
        f"Industry: {data.industry}; financial={data.financial}."
        if data.industry else "Industry unavailable; nonfinancial assumed."))
    _header(sheet, 18, ["Forecast", "Year 1", "Year 2", "Year 3", "Year 4",
                        "Year 5"])
    labels = {
        19: "Forecast end", 20: "Revenue", 21: "EBIT margin",
        22: "EBIT", 23: "Operating cash tax", 24: "NOPAT",
        25: "D&A", 26: "Cash capex", 27: "Operating NWC",
        28: "Change in operating NWC", 29: "FCFF", 30: "Discount factor",
        31: "PV of FCFF", 32: "Extra reinvestment / cash use (+/-)",
        34: "WACC", 35: "Terminal growth",
        36: "Terminal FCFF", 37: "Terminal enterprise value",
        38: "PV of terminal value", 39: "Enterprise value",
        40: "Cash less debt plus other common-equity adjustments",
        41: "Common equity value", 42: "Common equity / shares",
    }
    for row, label in labels.items():
        sheet.cell(row, 2, label)
    for col in range(3, 8):
        letter = get_column_letter(col)
        previous = get_column_letter(col - 1)
        count = col - 2
        _formula(sheet, 19, col,
                 f"=DATE(YEAR($C$9)+{count},MONTH($C$9),DAY($C$9))")
        sheet.cell(19, col).number_format = "yyyy-mm-dd"
        base = "Assumptions!$C$7" if col == 3 else f"{previous}20"
        _formula(sheet, 20, col, (
            f'=IF(AND($C$16=1,COUNT({base},Assumptions!{letter}16)=2),'
            f'IF(AND({base}>0,Assumptions!{letter}16>-1),'
            f'{base}*(1+Assumptions!{letter}16),""),"")'))
        _formula(sheet, 21, col, (
            f'=IF(ISNUMBER(Assumptions!{letter}21),'
            f'Assumptions!{letter}21,"")'))
        _formula(sheet, 22, col, (
            f'=IF(COUNT({letter}20:{letter}21)=2,'
            f'{letter}20*{letter}21,"")'))
        _formula(sheet, 23, col, (
            f'=IF(COUNT({letter}22,Assumptions!{letter}26)=2,'
            f'IF(AND(Assumptions!{letter}26>=0,Assumptions!{letter}26<=1),'
            f'MAX({letter}22,0)*Assumptions!{letter}26,""),"")'))
        _formula(sheet, 24, col, (
            f'=IF(COUNT({letter}22:{letter}23)=2,'
            f'{letter}22-{letter}23,"")'))
        for row, driver in ((25, 31), (26, 36), (27, 41)):
            nonnegative = (f'Assumptions!{letter}{driver}>=0'
                           if row != 27 else "TRUE")
            _formula(sheet, row, col, (
                f'=IF(COUNT({letter}20,Assumptions!{letter}{driver})=2,'
                f'IF({nonnegative},'
                f'{letter}20*Assumptions!{letter}{driver},""),"")'))
        prior_nwc = "Assumptions!$C$8" if col == 3 else f"{previous}27"
        _formula(sheet, 28, col, (
            f'=IF(COUNT({letter}27,{prior_nwc})=2,'
            f'{letter}27-{prior_nwc},"")'))
        _formula(sheet, 32, col, (
            f'=IF(ISNUMBER(Assumptions!{letter}60),'
            f'Assumptions!{letter}60,"")'))
        _formula(sheet, 29, col, (
            f'=IF(COUNT({letter}24:{letter}26,{letter}28,{letter}32)=5,'
            f'{letter}24+{letter}25-{letter}26-{letter}28-{letter}32,"")'))
        _formula(sheet, 30, col, (
            f'=IF(ISNUMBER($C$34),IF($C$34>0,'
            f'1/(1+$C$34)^{count},""),"")'))
        _formula(sheet, 31, col, (
            f'=IF(COUNT({letter}29:{letter}30)=2,'
            f'{letter}29*{letter}30,"")'))
        sheet.cell(21, col).number_format = RATE
        sheet.cell(30, col).number_format = "0.000"
    _formula(sheet, 34, 3,
             '=IF(ISNUMBER(Assumptions!C56),Assumptions!C56,"")')
    _formula(sheet, 35, 3,
             '=IF(ISNUMBER(Assumptions!C57),Assumptions!C57,"")')
    _formula(sheet, 36, 3,
             '=IF(COUNT(C55:C56)=2,C55-C56,"")')
    _formula(sheet, 37, 3, (
        '=IF(COUNT(C34:C36)=3,IF(AND(C34>C35,C34>0,C34<=1),'
        'C36/(C34-C35),""),"")'))
    _formula(sheet, 38, 3,
             '=IF(COUNT(C37,G30)=2,C37*G30,"")')
    _formula(sheet, 39, 3,
             '=IF(COUNT(C31:G31,C38)=6,SUM(C31:G31,C38),"")')
    _formula(sheet, 40, 3, (
        '=IF(COUNT(Assumptions!C9:C10,Assumptions!C12)=3,'
        'IF(AND(Assumptions!C9>=0,Assumptions!C10>=0),'
        'Assumptions!C9-Assumptions!C10+Assumptions!C12,""),"")'))
    _formula(sheet, 41, 3,
             '=IF(COUNT(C39:C40)=2,SUM(C39:C40),"")')
    _formula(sheet, 42, 3, (
        '=IF(C16=1,IF(COUNT(C41,Assumptions!C11)=2,'
        'IF(Assumptions!C11>0,C41/Assumptions!C11,""),""),"")'))
    sheet["C34"].number_format = RATE
    sheet["C35"].number_format = RATE
    sheet["C42"].number_format = '0.00;(0.00);"-"'
    sheet["B46"] = "Sensitivity: value per common share (TWD)"
    sheet["I46"] = "Enter side values; center links to selected WACC / growth."
    sheet["B48"] = "WACC / g"
    for col in range(4, 7):
        _input(sheet, 48, col)
        sheet.cell(48, col).number_format = RATE
    _formula(sheet, 48, 5, '=IF(ISNUMBER(C35),C35,"")')
    for row in range(49, 52):
        _input(sheet, row, 3)
        sheet.cell(row, 3).number_format = RATE
        for col in range(4, 7):
            letter = get_column_letter(col)
            rate = f"$C{row}"
            growth = f"{letter}$48"
            pv = "+".join(
                f"{get_column_letter(index + 2)}29/(1+{rate})^{index}"
                for index in range(1, 6))
            _formula(sheet, row, col, (
                f'=IF(COUNT($C$42,{rate},{growth})=3,'
                f'IF(AND({rate}>{growth},{rate}>0,{rate}<=1,{growth}>-1),'
                f'({pv}+$G$24*(1+{growth})*(1-{growth}/$C$54)/'
                f'({rate}-{growth})/'
                f'(1+{rate})^5+$C$40)/Assumptions!$C$11,""),"")'))
            sheet.cell(row, col).number_format = '0.00;(0.00);"-"'
    _formula(sheet, 50, 3, '=IF(ISNUMBER(C34),C34,"")')
    for row, label in (
        (54, "Terminal ROIC"), (55, "Terminal NOPAT (Year 6)"),
        (56, "Terminal net reinvestment"),
        (57, "Year 5 net reinvestment incl. cash adjustments"),
        (58, "Terminal net reinvestment / NOPAT"),
        (59, "Year 5 net reinvestment / NOPAT"),
        (60, "Change vs Year 5 reinvestment grown by g"),
    ):
        sheet.cell(row, 2, label)
    _formula(sheet, 54, 3,
             '=IF(ISNUMBER(Assumptions!C58),Assumptions!C58,"")')
    _formula(sheet, 55, 3,
             '=IF(COUNT(G24,C35)=2,IF(C35>-1,G24*(1+C35),""),"")')
    _formula(sheet, 56, 3, (
        '=IF(COUNT(C54:C55,C35)=3,IF(C54>0,C55*C35/C54,""),"")'))
    _formula(sheet, 57, 3, (
        '=IF(COUNT(G25:G26,G28,G32)=4,G26-G25+G28+G32,"")'))
    _formula(sheet, 58, 3,
             '=IF(COUNT(C55:C56)=2,IF(C55<>0,C56/C55,""),"")')
    _formula(sheet, 59, 3,
             '=IF(COUNT(C57,G24)=2,IF(G24<>0,C57/G24,""),"")')
    _formula(sheet, 60, 3,
             '=IF(COUNT(C56:C57,C35)=3,C56-C57*(1+C35),"")')
    for row in (54, 58, 59):
        sheet.cell(row, 3).number_format = RATE
    sheet["I57"] = (
        "Includes extra after-tax cash use; avoid capex/NWC overlap.")
    sheet["I60"] = (
        "Explain the transition in research notes; no approval checkbox.")
    sheet.sheet_properties.tabColor = "203548"


def write_workbook(
    data: Dataset, path: Path, supplement: dict[str, Any] | None = None
) -> dict[str, Any]:
    """Write one formula workbook; missing inputs remain editable and gated.

    Supplied supplements must first pass load_supplement(). The output path
    must not exist, preventing accidental overwrite. Returns intake coverage,
    not a Python-computed valuation. Raises on conflicting annual facts.
    """
    if path.exists():
        raise FileExistsError(path)
    supplement = supplement or {}
    facts = data.facts + [Fact(**row) for row in supplement.get("facts", [])]
    _validate_facts(facts)
    for values in [supplement.get("inputs", {}), *supplement.get(
            "scenarios", {}).values()]:
        for item in values.values():
            _evidenced(item)
    book = Workbook()
    book.remove(book.active)
    for title in ("Model", "Assumptions", "History", "Facts", "Evidence"):
        book.create_sheet(title)
    book.calculation = CalcProperties(
        calcId=191029, fullCalcOnLoad=True, forceFullCalc=True,
        calcMode="auto")
    for title in ("Model", "Assumptions", "History"):
        # Defaults precede number formats and color overrides below.
        for row in book[title].iter_rows(min_row=1, max_row=70, max_col=11):
            for cell in row:
                cell.font = Font(name="Arial", size=10, color=INK)
                cell.alignment = Alignment(vertical="center")
        _style(book[title], 11)
    _history(book, data, facts)
    _facts(book, facts, data.gaps)
    _assumptions(book, data, facts, supplement)
    _model(book, data)
    _research(book, supplement)
    for title in ("Model", "Assumptions", "History"):
        book[title]["B2"].font = Font(name="Arial", size=14, bold=True)
        book[title].print_options.horizontalCentered = True
        book[title].sheet_properties.pageSetUpPr.fitToPage = True
        book[title].page_setup.orientation = "landscape"
        book[title].page_setup.paperSize = "9"
        book[title].page_setup.fitToWidth = 1
        book[title].page_setup.fitToHeight = 0
    book["Assumptions"].column_dimensions["J"].width = 70
    book["Assumptions"].column_dimensions["K"].width = 70
    book["Facts"].column_dimensions["A"].width = 26
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    coverage = {str(year): sum(annual_value(facts, metric, year) is not None
                               for metric in HISTORY_METRICS)
                for year in range(data.valuation_date.year - 5,
                                  data.valuation_date.year)}
    return {"path": str(path.resolve()), "ticker": data.ticker,
            "financial": data.financial, "annual_metric_counts": coverage,
            "fact_count": len(facts), "gaps": data.gaps,
            "recalculated": False}


def create_current_workbook(
    ticker: str, path: Path, supplement_path: Path | None = None,
    capture_dir: Path | None = None,
) -> dict[str, Any]:
    """Collect current official data and create a useful incomplete workbook.

    No credentials or paid providers are used. Input supplements are local
    JSON. This function does not run an office application or evaluate Excel.
    """
    supplement = load_supplement(supplement_path) if supplement_path else {}
    data = collect_current_data(ticker, capture_dir)
    return write_workbook(data, path, supplement)


def recalculate_workbook(
    path: Path, executable: str = "soffice", timeout: float = 120
) -> dict[str, Any]:
    """Recalculate with installed LibreOffice and record native readback.

    Uses an isolated temporary office profile. On success, atomically replace
    the supplied workbook and save a sibling .recalc.json record. Missing
    software, conversion errors or formula errors raise without replacement.
    This is an explicit local write, not an automatic software installer.
    """
    command = shutil.which(executable)
    if command is None:
        raise FileNotFoundError("LibreOffice executable not available")
    path = path.resolve(strict=True)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    version = subprocess.run(
        [command, "--headless", "--version"], check=True,
        capture_output=True, text=True, timeout=20).stdout.strip()
    with tempfile.TemporaryDirectory(prefix="taiwan-dcf-") as temporary:
        folder = Path(temporary)
        result = subprocess.run([
            command, f"-env:UserInstallation={(folder / 'profile').as_uri()}",
            "--headless", "--convert-to", "xlsx", "--outdir", str(folder),
            str(path),
        ], check=True, capture_output=True, text=True, timeout=timeout)
        converted = folder / path.name
        if not converted.is_file():
            raise RuntimeError(f"Native conversion did not create a file: "
                               f"{result.stderr.strip()}")
        from openpyxl import load_workbook

        cached = load_workbook(converted, data_only=True)
        formulas = load_workbook(converted, data_only=False)
        errors = [f"{sheet.title}!{cell.coordinate}" for sheet in cached
                  for row in sheet.iter_rows() for cell in row
                  if cell.data_type == "e"]
        if errors or formulas["Model"]["C42"].data_type != "f":
            raise ValueError(f"Native formula check failed: {errors[:10]}")
        price = cached["Model"]["C6"].value
        record = {
            "engine": "LibreOffice", "engine_version": version,
            "recalculated_at": datetime.now(timezone.utc).isoformat(),
            "input_sha256": before,
            "output_sha256": hashlib.sha256(
                converted.read_bytes()).hexdigest(),
            "valuation_available": isinstance(price, (int, float)),
            "common_equity_per_share": price,
            "status": cached["Model"]["C7"].value,
            "evidence_grade": cached["Model"]["C12"].value,
            "formula_error_count": 0,
        }
        # Stage beside the destination so replacement stays on one filesystem.
        with tempfile.NamedTemporaryFile(
                dir=path.parent, prefix=".dcf-recalc-",
                delete=False) as handle:
            staging = Path(handle.name)
        try:
            shutil.copyfile(converted, staging)
            staging.replace(path)
        finally:
            staging.unlink(missing_ok=True)
    record_path = path.with_suffix(".recalc.json")
    record_path.write_text(json.dumps(record, ensure_ascii=False, indent=2)
                           + "\n", encoding="utf-8")
    return record


def create_researched_workbook(
    data: Dataset | Path, supplement_path: Path, path: Path,
    recalculate: bool = False,
) -> dict[str, Any]:
    """Build offline from API captures or browser-researched local JSON.

    No network access is needed. The Dataset JSON may come from
    collect_current_data or an agent's source-backed browser research.
    This function does not research or invent missing assumptions. Optional
    native recalculation requires an installed office engine.
    """
    if isinstance(data, Path):
        raw = json.loads(data.read_text(encoding="utf-8"))
        raw["valuation_date"] = date.fromisoformat(raw["valuation_date"])
        raw["facts"] = [Fact(**row) for row in raw.get("facts", [])]
        data = Dataset(**raw)
    result = write_workbook(data, path, load_supplement(supplement_path))
    if recalculate:
        result["native"] = recalculate_workbook(path)
        result["recalculated"] = True
    return result
